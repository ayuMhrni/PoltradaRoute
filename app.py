from flask import Flask, render_template, request, redirect, url_for, send_file, flash
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path
from datetime import datetime
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import json
import math
import tempfile
import os
import time
import copy

app = Flask(__name__)
app.secret_key = "poltrada-optimasi-rute-2026"

BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_EXCEL = BASE_DIR / "IDEA FEST6(1).xlsx"

if not TEMPLATE_EXCEL.exists():
    alt = BASE_DIR / "IDEA FEST6.xlsx"
    if alt.exists():
        TEMPLATE_EXCEL = alt


# ============================================================
# DATA AKTIF
# Data pengguna disimpan sementara selama aplikasi berjalan.
# Excel hanya digunakan sebagai template/contoh laporan.
# ============================================================

ACTIVE = {
    "gudang": {
        "nama": "",
        "alamat": "",
        "lat": None,
        "lon": None
    },
    "outlet": [],
    "jumlah_kendaraan": 3,
    "kapasitas": 5000,
    "harga_bbm": 6800,
    "biaya_supir": 25000,
    "hasil": None,
    "api_status": None,
    "last_process": None,
}

HISTORY = []
REDO_STACK = []

# Konsumsi internal untuk perhitungan biaya BBM.
# Tidak ditampilkan sebagai pengaturan pengguna.
BBM_KM_PER_LITER = 5.0


# ============================================================
# UTILITAS
# ============================================================

def now_text():
    return datetime.now().strftime("%d/%m/%Y %H:%M:%S")


def clean(v):
    return "" if v is None else str(v).strip()


def num(v, default=0):
    try:
        return float(v)
    except Exception:
        return default


def snapshot():
    return copy.deepcopy(ACTIVE)


def restore(data):
    ACTIVE.clear()
    ACTIVE.update(copy.deepcopy(data))


def save_history():
    HISTORY.append(snapshot())

    if len(HISTORY) > 30:
        HISTORY.pop(0)

    REDO_STACK.clear()


def undo():
    if not HISTORY:
        return False

    REDO_STACK.append(snapshot())
    restore(HISTORY.pop())

    return True


def redo():
    if not REDO_STACK:
        return False

    HISTORY.append(snapshot())
    restore(REDO_STACK.pop())

    return True


# ============================================================
# API
# ============================================================

def get_json(url, timeout=60):
    req = Request(
        url,
        headers={
            "User-Agent": "Poltrada-Bali-Sistem-Optimasi/3.0"
        }
    )

    with urlopen(req, timeout=timeout) as response:
        return json.loads(
            response.read().decode("utf-8")
        )


def geocode(address):
    params = urlencode({
        "q": address,
        "format": "json",
        "limit": 1,
        "countrycodes": "id",
        "addressdetails": 1,
    })

    data = get_json(
        "https://nominatim.openstreetmap.org/search?"
        + params,
        45
    )

    if not data:
        raise ValueError(
            f"Koordinat tidak ditemukan untuk alamat: {address}"
        )

    return (
        float(data[0]["lat"]),
        float(data[0]["lon"])
    )


def osrm_table(points):
    coords = ";".join(
        f"{lon},{lat}"
        for lat, lon in points
    )

    url = (
        f"https://router.project-osrm.org/table/v1/driving/"
        f"{coords}"
        "?annotations=distance,duration"
    )

    result = get_json(url, 90)

    if result.get("code") != "Ok":
        raise ValueError(
            "OSRM tidak berhasil mengembalikan matriks."
        )

    return (
        result["distances"],
        result["durations"]
    )


def calculate_api_matrix():
    if not ACTIVE["gudang"]["alamat"]:
        raise ValueError(
            "Alamat gudang belum diisi."
        )

    if not ACTIVE["outlet"]:
        raise ValueError(
            "Belum ada data outlet."
        )

    # Geocoding gudang
    glat, glon = geocode(
        ACTIVE["gudang"]["alamat"]
    )

    ACTIVE["gudang"]["lat"] = glat
    ACTIVE["gudang"]["lon"] = glon

    # Geocoding outlet
    for outlet in ACTIVE["outlet"]:
        lat, lon = geocode(
            outlet["alamat"]
        )

        outlet["lat"] = lat
        outlet["lon"] = lon

        # Memberi jeda agar penggunaan Nominatim tetap wajar.
        time.sleep(1)

    labels = (
        ["GUDANG"]
        + [o["kode"] for o in ACTIVE["outlet"]]
    )

    points = [
        (
            ACTIVE["gudang"]["lat"],
            ACTIVE["gudang"]["lon"]
        )
    ]

    points += [
        (
            o["lat"],
            o["lon"]
        )
        for o in ACTIVE["outlet"]
    ]

    distances, durations = osrm_table(points)

    matrix_distance = {}
    matrix_time = {}

    for i, a in enumerate(labels):
        matrix_distance[a] = {}
        matrix_time[a] = {}

        for j, b in enumerate(labels):
            matrix_distance[a][b] = round(
                distances[i][j] / 1000,
                3
            )

            matrix_time[a][b] = round(
                durations[i][j] / 60,
                1
            )

    ACTIVE["api_status"] = "Nominatim + OSRM"
    ACTIVE["last_process"] = now_text()

    return (
        labels,
        matrix_distance,
        matrix_time
    )


# ============================================================
# VRP + NEAREST NEIGHBOR
# ============================================================

def route_distance(route, matrix):
    return sum(
        matrix[route[i]][route[i + 1]]
        for i in range(len(route) - 1)
    )


def nearest_neighbor(customers, matrix):
    remaining = set(customers)

    current = "GUDANG"
    route = ["GUDANG"]

    while remaining:
        nxt = min(
            remaining,
            key=lambda x: (
                matrix[current][x],
                x
            )
        )

        route.append(nxt)
        remaining.remove(nxt)
        current = nxt

    route.append("GUDANG")

    return route


def sweep_order(outlets):
    depot = ACTIVE["gudang"]

    dlat = depot["lat"]
    dlon = depot["lon"]

    def angle(o):
        return math.atan2(
            o["lat"] - dlat,
            o["lon"] - dlon
        )

    return sorted(
        outlets,
        key=lambda o: (
            angle(o),
            o["kode"]
        )
    )


def vrp_allocate(outlets):
    k = int(ACTIVE["jumlah_kendaraan"])
    capacity = float(ACTIVE["kapasitas"])

    total = sum(
        float(o["permintaan"])
        for o in outlets
    )

    if total > k * capacity:
        raise ValueError(
            f"Total permintaan {total:,.0f} kg "
            f"melebihi kapasitas total "
            f"{k * capacity:,.0f} kg."
        )

    ordered = sweep_order(outlets)
    n = len(ordered)

    demand = {
        o["kode"]: float(o["permintaan"])
        for o in outlets
    }

    if k > n:
        k = n

    best = None

    if k <= 5 and n <= 40:

        def partitions(seq, parts, start=0):

            if parts == 1:
                yield [seq[start:]]
                return

            for cut in range(
                start + 1,
                n - parts + 2
            ):
                for rest in partitions(
                    seq,
                    parts - 1,
                    cut
                ):
                    yield [
                        seq[start:cut]
                    ] + rest

        for rotation in range(n):

            seq = (
                ordered[rotation:]
                + ordered[:rotation]
            )

            for groups in partitions(
                seq,
                k
            ):

                loads = [
                    sum(
                        demand[o["kode"]]
                        for o in group
                    )
                    for group in groups
                ]

                if any(
                    load > capacity
                    for load in loads
                ):
                    continue

                avg = total / k

                balance = sum(
                    (load - avg) ** 2
                    for load in loads
                ) / max(
                    avg ** 2,
                    1
                )

                if (
                    best is None
                    or balance < best[0]
                ):
                    best = (
                        balance,
                        groups
                    )

    if best is None:

        groups = [
            []
            for _ in range(k)
        ]

        loads = [0] * k

        for o in ordered:

            q = demand[o["kode"]]

            possible = [
                i
                for i in range(k)
                if loads[i] + q <= capacity
            ]

            if not possible:
                raise ValueError(
                    "Tidak ditemukan pembagian VRP "
                    "yang memenuhi kapasitas."
                )

            idx = min(
                possible,
                key=lambda i: loads[i]
            )

            groups[idx].append(o)
            loads[idx] += q

    else:
        groups = best[1]

    return groups


def build_result(
    matrix_distance,
    matrix_time
):
    outlets = ACTIVE["outlet"]

    groups = vrp_allocate(outlets)

    # --------------------------------------------------------
    # PEMBANDING AWAL
    # Urutan input pengguna dibagi sesuai kapasitas.
    # --------------------------------------------------------

    baseline_groups = []

    current = []
    load = 0

    for o in outlets:

        q = float(o["permintaan"])

        if (
            current
            and load + q > ACTIVE["kapasitas"]
        ):
            baseline_groups.append(current)

            current = []
            load = 0

        current.append(o)
        load += q

    if current:
        baseline_groups.append(current)

    baseline_distance = 0

    for group in baseline_groups:

        route = (
            ["GUDANG"]
            + [o["kode"] for o in group]
            + ["GUDANG"]
        )

        baseline_distance += route_distance(
            route,
            matrix_distance
        )

    # --------------------------------------------------------
    # RUTE HASIL OPTIMASI
    # --------------------------------------------------------

    routes = []

    for idx, group in enumerate(
        groups,
        1
    ):

        codes = [
            o["kode"]
            for o in group
        ]

        route = nearest_neighbor(
            codes,
            matrix_distance
        )

        distance = route_distance(
            route,
            matrix_distance
        )

        travel_time = route_distance(
            route,
            matrix_time
        )

        load = sum(
            float(o["permintaan"])
            for o in group
        )

        fuel_liter = (
            distance
            / BBM_KM_PER_LITER
        )

        fuel_cost = (
            fuel_liter
            * ACTIVE["harga_bbm"]
        )

        driver_cost = (
            travel_time / 60
        ) * ACTIVE["biaya_supir"]

        operational_cost = (
            fuel_cost
            + driver_cost
        )

        routes.append({
            "kendaraan": idx,
            "jenis": "Colt Diesel Double",
            "route": route,
            "muatan": load,
            "utilisasi": (
                load
                / ACTIVE["kapasitas"]
                * 100
            ),
            "jarak": distance,
            "waktu": travel_time,
            "bbm": fuel_liter,
            "biaya_bbm": fuel_cost,
            "biaya_supir": driver_cost,
            "biaya": operational_cost,
        })

    optimized = sum(
        r["jarak"]
        for r in routes
    )

    saving = (
        baseline_distance
        - optimized
    )

    percent = (
        saving
        / baseline_distance
        * 100
        if baseline_distance
        else 0
    )

    total_cost = sum(
        r["biaya"]
        for r in routes
    )

    result = {
        "jumlah_outlet": len(outlets),

        "total_muatan": sum(
            float(o["permintaan"])
            for o in outlets
        ),

        "jarak_awal": baseline_distance,

        "jarak_optimasi": optimized,

        "penghematan": saving,

        "persen": percent,

        "total_waktu": sum(
            r["waktu"]
            for r in routes
        ),

        "total_biaya": total_cost,

        "rute": routes,

        "timestamp": now_text(),

        "api_status": ACTIVE["api_status"],

        "harga_bbm": ACTIVE["harga_bbm"],

        "biaya_supir": ACTIVE["biaya_supir"],

        "metode": (
            "VRP (pembagian berdasarkan kapasitas) "
            "→ Nearest Neighbor "
            "(urutan kunjungan)"
        ),
    }

    ACTIVE["hasil"] = result

    ACTIVE["last_process"] = (
        result["timestamp"]
    )

    return result


# ============================================================
# EXCEL REKAP
# ============================================================

def style_sheet(ws):
    blue = "0759A6"

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=blue
        )

        cell.font = Font(
            color="FFFFFF",
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


def make_download_file():

    if TEMPLATE_EXCEL.exists():
        wb = load_workbook(
            TEMPLATE_EXCEL
        )
    else:
        from openpyxl import Workbook
        wb = Workbook()

    for name in [
        "REKAP WEB",
        "DATA INPUT WEB",
        "HASIL OPTIMASI WEB",
    ]:

        if name in wb.sheetnames:
            del wb[name]

    # --------------------------------------------------------
    # REKAP WEB
    # --------------------------------------------------------

    info = wb.create_sheet(
        "REKAP WEB"
    )

    info.append([
        "INFORMASI SISTEM",
        "NILAI"
    ])

    info.append([
        "Tanggal & Jam Download",
        now_text()
    ])

    info.append([
        "Status Sistem",
        "Berhasil diproses"
    ])

    info.append([
        "Sumber Geocoding",
        "Nominatim / OpenStreetMap"
    ])

    info.append([
        "Sumber Routing",
        "OSRM"
    ])

    info.append([
        "Metode Optimasi",
        (
            ACTIVE["hasil"]["metode"]
            if ACTIVE["hasil"]
            else "Belum ada"
        )
    ])

    info.append([
        "Jumlah Kendaraan",
        ACTIVE["jumlah_kendaraan"]
    ])

    info.append([
        "Kapasitas Kendaraan (kg)",
        ACTIVE["kapasitas"]
    ])

    info.append([
        "Harga BBM (Rp/liter)",
        ACTIVE["harga_bbm"]
    ])

    info.append([
        "Biaya Supir (Rp/jam)",
        ACTIVE["biaya_supir"]
    ])

    info.append([
        "Jumlah Outlet",
        len(ACTIVE["outlet"])
    ])

    info.append([
        "Waktu Proses Terakhir",
        ACTIVE["last_process"] or "-"
    ])

    style_sheet(info)

    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 75

    # --------------------------------------------------------
    # DATA INPUT WEB
    # --------------------------------------------------------

    data_ws = wb.create_sheet(
        "DATA INPUT WEB"
    )

    data_ws.append([
        "KODE",
        "NAMA OUTLET",
        "ALAMAT",
        "PERMINTAAN (KG)",
        "LATITUDE",
        "LONGITUDE"
    ])

    for o in ACTIVE["outlet"]:

        data_ws.append([
            o["kode"],
            o["nama"],
            o["alamat"],
            o["permintaan"],
            o["lat"],
            o["lon"]
        ])

    style_sheet(data_ws)

    for c, w in {
        "A": 10,
        "B": 28,
        "C": 65,
        "D": 18,
        "E": 16,
        "F": 16
    }.items():

        data_ws.column_dimensions[c].width = w

    # --------------------------------------------------------
    # HASIL OPTIMASI WEB
    # --------------------------------------------------------

    if ACTIVE["hasil"]:

        hasil = ACTIVE["hasil"]

        hws = wb.create_sheet(
            "HASIL OPTIMASI WEB"
        )

        hws.append([
            "PARAMETER",
            "HASIL"
        ])

        for row in [

            [
                "Jarak Awal (km)",
                round(
                    hasil["jarak_awal"],
                    2
                )
            ],

            [
                "Jarak Optimasi (km)",
                round(
                    hasil["jarak_optimasi"],
                    2
                )
            ],

            [
                "Penghematan Jarak (km)",
                round(
                    hasil["penghematan"],
                    2
                )
            ],

            [
                "Penghematan (%)",
                round(
                    hasil["persen"],
                    2
                )
            ],

            [
                "Total Muatan (kg)",
                round(
                    hasil["total_muatan"],
                    2
                )
            ],

            [
                "Total Waktu (menit)",
                round(
                    hasil["total_waktu"],
                    1
                )
            ],

            [
                "Total Biaya (Rp)",
                round(
                    hasil["total_biaya"],
                    0
                )
            ],

            [
                "Harga BBM (Rp/liter)",
                hasil["harga_bbm"]
            ],

            [
                "Biaya Supir (Rp/jam)",
                hasil["biaya_supir"]
            ],

            [
                "Tanggal & Jam Hasil",
                hasil["timestamp"]
            ],

        ]:

            hws.append(row)

        hws.append([])

        hws.append([
            "KENDARAAN",
            "JENIS",
            "RUTE",
            "MUATAN (KG)",
            "UTILISASI (%)",
            "JARAK (KM)",
            "WAKTU (MENIT)",
            "BBM (L)",
            "BIAYA BBM (Rp)",
            "BIAYA SUPIR (Rp)",
            "BIAYA OPERASIONAL (Rp)"
        ])

        for r in hasil["rute"]:

            hws.append([
                r["kendaraan"],
                r["jenis"],
                " → ".join(
                    r["route"]
                ),
                r["muatan"],
                round(
                    r["utilisasi"],
                    2
                ),
                round(
                    r["jarak"],
                    3
                ),
                round(
                    r["waktu"],
                    1
                ),
                round(
                    r["bbm"],
                    2
                ),
                round(
                    r["biaya_bbm"],
                    0
                ),
                round(
                    r["biaya_supir"],
                    0
                ),
                round(
                    r["biaya"],
                    0
                ),
            ])

        style_sheet(hws)

        for c, w in {
            "A": 14,
            "B": 24,
            "C": 75,
            "D": 16,
            "E": 16,
            "F": 16,
            "G": 18,
            "H": 12,
            "I": 20,
            "J": 20,
            "K": 24
        }.items():

            hws.column_dimensions[c].width = w

    fd, path = tempfile.mkstemp(
        suffix=".xlsx"
    )

    os.close(fd)

    wb.save(path)

    return path


# ============================================================
# ROUTES
# ============================================================

@app.route("/")
def index():

    return render_template(
        "index.html",
        active=ACTIVE,
        now=now_text(),
        can_undo=bool(HISTORY),
        can_redo=bool(REDO_STACK),
    )


@app.post("/simpan-gudang")
def simpan_gudang():

    save_history()

    ACTIVE["gudang"] = {
        "nama": clean(
            request.form.get("nama")
        ),
        "alamat": clean(
            request.form.get("alamat")
        ),
        "lat": None,
        "lon": None,
    }

    ACTIVE["hasil"] = None
    ACTIVE["api_status"] = None

    flash(
        "Data gudang berhasil diperbarui.",
        "success"
    )

    return redirect(
        url_for("index") + "#gudang"
    )


@app.post("/simpan-armada")
def simpan_armada():

    jumlah = int(
        request.form.get(
            "jumlah_kendaraan",
            3
        )
    )

    kapasitas = float(
        request.form.get(
            "kapasitas",
            5000
        )
    )

    if jumlah < 1 or kapasitas <= 0:

        flash(
            "Jumlah kendaraan dan kapasitas harus lebih dari 0.",
            "error"
        )

    else:

        save_history()

        ACTIVE["jumlah_kendaraan"] = jumlah
        ACTIVE["kapasitas"] = kapasitas
        ACTIVE["hasil"] = None

        flash(
            "Konfigurasi armada berhasil diperbarui.",
            "success"
        )

    return redirect(
        url_for("index") + "#armada"
    )


@app.post("/simpan-pengaturan")
def simpan_pengaturan():

    harga_bbm = num(
        request.form.get(
            "harga_bbm"
        ),
        6800
    )

    biaya_supir = num(
        request.form.get(
            "biaya_supir"
        ),
        25000
    )

    if harga_bbm <= 0 or biaya_supir < 0:

        flash(
            "Harga BBM harus lebih dari 0 dan biaya supir tidak boleh negatif.",
            "error"
        )

    else:

        save_history()

        ACTIVE["harga_bbm"] = harga_bbm
        ACTIVE["biaya_supir"] = biaya_supir
        ACTIVE["hasil"] = None

        flash(
            "Pengaturan biaya berhasil diperbarui.",
            "success"
        )

    return redirect(
        url_for("index") + "#pengaturan"
    )


@app.post("/simpan-outlet")
def simpan_outlet():

    kode = clean(
        request.form.get("kode")
    ).upper()

    nama = clean(
        request.form.get("nama")
    )

    alamat = clean(
        request.form.get("alamat")
    )

    permintaan = num(
        request.form.get(
            "permintaan"
        )
    )

    if (
        not kode
        or not nama
        or not alamat
        or permintaan <= 0
    ):

        flash(
            "Kode, nama, alamat, dan permintaan wajib diisi.",
            "error"
        )

        return redirect(
            url_for("index") + "#outlet"
        )

    save_history()

    existing = next(
        (
            o
            for o in ACTIVE["outlet"]
            if o["kode"] == kode
        ),
        None
    )

    if existing:

        existing.update({
            "nama": nama,
            "alamat": alamat,
            "permintaan": permintaan,
            "lat": None,
            "lon": None,
        })

        pesan = (
            f"Outlet {kode} berhasil diperbarui."
        )

    else:

        ACTIVE["outlet"].append({
            "kode": kode,
            "nama": nama,
            "alamat": alamat,
            "permintaan": permintaan,
            "lat": None,
            "lon": None,
        })

        pesan = (
            f"Outlet {kode} berhasil ditambahkan."
        )

    ACTIVE["outlet"].sort(
        key=lambda x: x["kode"]
    )

    ACTIVE["hasil"] = None
    ACTIVE["api_status"] = None

    flash(
        pesan,
        "success"
    )

    return redirect(
        url_for("index") + "#outlet"
    )


@app.get("/hapus/<kode>")
def hapus_outlet(kode):

    kode = kode.upper()

    existing = next(
        (
            o
            for o in ACTIVE["outlet"]
            if o["kode"] == kode
        ),
        None
    )

    if existing is None:

        flash(
            f"Outlet {kode} tidak ditemukan.",
            "error"
        )

    else:

        save_history()

        ACTIVE["outlet"] = [
            o
            for o in ACTIVE["outlet"]
            if o["kode"] != kode
        ]

        ACTIVE["hasil"] = None
        ACTIVE["api_status"] = None

        flash(
            f"Outlet {kode} berhasil dihapus.",
            "success"
        )

    return redirect(
        url_for("index") + "#outlet"
    )


@app.post("/hapus-semua-data")
def hapus_semua_data():

    save_history()

    ACTIVE["gudang"] = {
        "nama": "",
        "alamat": "",
        "lat": None,
        "lon": None,
    }

    ACTIVE["outlet"] = []

    ACTIVE["jumlah_kendaraan"] = 3
    ACTIVE["kapasitas"] = 5000
    ACTIVE["harga_bbm"] = 6800
    ACTIVE["biaya_supir"] = 25000

    ACTIVE["hasil"] = None
    ACTIVE["api_status"] = None
    ACTIVE["last_process"] = None

    flash(
        "Semua data input aktif berhasil dihapus.",
        "success"
    )

    return redirect(
        url_for("index") + "#reset"
    )


@app.post("/undo")
def undo_action():

    if undo():

        flash(
            "Perubahan terakhir berhasil dibatalkan.",
            "success"
        )

    else:

        flash(
            "Belum ada perubahan yang dapat di-undo.",
            "error"
        )

    return redirect(
        url_for("index") + "#outlet"
    )


@app.post("/redo")
def redo_action():

    if redo():

        flash(
            "Perubahan yang dibatalkan berhasil dikembalikan.",
            "success"
        )

    else:

        flash(
            "Belum ada perubahan yang dapat di-redo.",
            "error"
        )

    return redirect(
        url_for("index") + "#outlet"
    )


@app.post("/hitung")
def hitung():

    try:

        save_history()

        calculate_api_matrix()

        flash(
            "Koordinat, jarak, dan waktu berhasil dihitung menggunakan API.",
            "success"
        )

    except Exception as e:

        if HISTORY:
            restore(
                HISTORY.pop()
            )

        flash(
            f"Perhitungan API gagal: {e}",
            "error"
        )

    return redirect(
        url_for("index") + "#hitung"
    )


@app.post("/optimasi")
def optimasi():

    try:

        if not ACTIVE["gudang"]["alamat"]:
            raise ValueError(
                "Alamat gudang belum diisi."
            )

        if not ACTIVE["outlet"]:
            raise ValueError(
                "Data outlet belum diisi."
            )

        save_history()

        _, matrix_distance, matrix_time = (
            calculate_api_matrix()
        )

        build_result(
            matrix_distance,
            matrix_time
        )

        flash(
            "Optimasi berhasil: VRP → Nearest Neighbor.",
            "success"
        )

        return render_template(
            "index.html",
            active=ACTIVE,
            now=now_text(),
            can_undo=bool(HISTORY),
            can_redo=bool(REDO_STACK),
            auto_result=True,
        )

    except Exception as e:

        if HISTORY:
            restore(
                HISTORY.pop()
            )

        flash(
            f"Optimasi gagal: {e}",
            "error"
        )

        return redirect(
            url_for("index") + "#optimasi"
        )


@app.get("/download")
def download():

    if not ACTIVE["outlet"]:

        flash(
            "Belum ada data input yang dapat direkap.",
            "error"
        )

        return redirect(
            url_for("index") + "#unduh"
        )

    try:

        path = make_download_file()

        return send_file(
            path,
            as_attachment=True,
            download_name=(
                f"REKAP_OPTIMASI_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            ),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    except Exception as e:

        flash(
            f"Download gagal: {e}",
            "error"
        )

        return redirect(
            url_for("index") + "#unduh"
        )


# ============================================================
# MENJALANKAN APLIKASI
# ============================================================

if __name__ == "__main__":

    print("=" * 65)
    print(
        "SISTEM OPTIMASI RUTE DISTRIBUSI - POLTRADA BALI"
    )
    print(
        "Data input aktif hanya disimpan sementara di memori."
    )
    print("=" * 65)

    # Jika dijalankan di SnapDeploy/server,
    # gunakan PORT yang diberikan oleh server.
    # Jika dijalankan di laptop,
    # gunakan port 5000 sebagai default.

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=False,
        use_reloader=False
    )